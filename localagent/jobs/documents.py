"""Text extraction for research sources (txt, md, PDF, DOCX), with a cache.

Downloaded and user-provided documents are untrusted input: extraction only, nothing embedded is executed, and
there are size limits. The cached text (with page markers for PDFs) is what quotes are verified against.
"""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

TEXT_TYPES = {".txt", ".md", ".markdown", ".rst", ".csv"}
DOC_TYPES = TEXT_TYPES | {".pdf", ".docx"}
MAX_FILE_BYTES = 80 * 2**20
MAX_CHARS = 3_000_000


@dataclass
class ExtractedText:
    path: Path
    text: str
    pages: int | None
    warning: str | None = None


def is_document(path: Path) -> bool:
    return path.suffix.lower() in DOC_TYPES


def normalize_ws(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


def _cache_path(cache_dir: Path, path: Path) -> Path:
    st = path.stat()
    digest = hashlib.sha1(f"{path.resolve()}|{st.st_size}|{st.st_mtime_ns}".encode()).hexdigest()[:16]
    return cache_dir / f"{path.stem[:60]}-{digest}.txt"


def extract(path: Path, cache_dir: Path | None = None) -> ExtractedText:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(str(path))
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError(f"{path.name} is larger than {MAX_FILE_BYTES // 2**20} MB")
    if cache_dir is not None:
        cached = _cache_path(cache_dir, path)
        if cached.exists():
            text = cached.read_text(encoding="utf-8")
            pages = text.count("\n[page ") if path.suffix.lower() == ".pdf" else None
            return ExtractedText(path, text, pages)
    suffix = path.suffix.lower()
    warning = None
    pages = None
    if suffix in TEXT_TYPES:
        text = path.read_text(encoding="utf-8", errors="replace")
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        parts = []
        for i, page in enumerate(reader.pages, 1):
            try:
                parts.append(f"\n[page {i}]\n{page.extract_text() or ''}")
            except Exception as e:           # damaged page: keep going
                parts.append(f"\n[page {i}]\n(could not extract: {type(e).__name__})")
        pages = len(reader.pages)
        text = "".join(parts)
        if len(normalize_ws(text.replace("[page", ""))) < 200 * max(1, pages // 4):
            warning = "Very little text was extracted; this PDF may be scanned images (OCR isn't supported yet)."
    elif suffix == ".docx":
        import docx
        d = docx.Document(str(path))
        blocks = []
        for p in d.paragraphs:
            if p.text.strip():
                style = (p.style.name or "").lower() if p.style is not None else ""
                blocks.append(("#" * int(style[-1]) + " " if style.startswith("heading") and style[-1:].isdigit() else "")
                              + p.text)
        for table in d.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(blocks)
    else:
        raise ValueError(f"Unsupported document type: {suffix}")
    text = text[:MAX_CHARS]
    if cache_dir is not None:
        cache_dir.mkdir(parents=True, exist_ok=True)
        _cache_path(cache_dir, path).write_text(text, encoding="utf-8")
    return ExtractedText(path, text, pages, warning)


def find_quote(text: str, quote: str) -> bool:
    """Whitespace- and case-insensitive containment, also tolerant of PDF hyphenation at line breaks."""
    q = normalize_ws(quote).lower()
    if len(q) < 8:
        return False
    t = normalize_ws(text).lower()
    if q in t:
        return True
    dehyphenated = normalize_ws(re.sub(r"(\w)-\s*\n\s*(\w)", r"\1\2", text)).lower()
    return q in dehyphenated


def closest_snippet(text: str, quote: str, width: int = 160) -> str | None:
    """Best-effort pointer for a failed quote: the passage sharing the most words with it."""
    words = [w for w in re.findall(r"\w+", quote.lower()) if len(w) > 3]
    if not words:
        return None
    flat = normalize_ws(text)
    low = flat.lower()
    best, best_score = None, 0
    step = max(40, width // 2)
    for i in range(0, max(1, len(low) - width), step):
        window = low[i:i + width * 2]
        score = sum(1 for w in set(words) if w in window)
        if score > best_score:
            best, best_score = flat[i:i + width * 2], score
    return best if best_score >= max(2, len(set(words)) // 3) else None
