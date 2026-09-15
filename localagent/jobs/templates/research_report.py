"""research_report: search documents, take verified notes, outline (user gate), write sections, compile (design §6.1)."""
from __future__ import annotations

import re
from pathlib import Path

from ..documents import DOC_TYPES
from . import register
from .base import HandlerResult, Template, is_approval

SKIP_DIRS = {"jobs", "sections", ".git", "__pycache__", "node_modules", ".venv"}
OUTPUT_NAMES = {"outline.md", "report.md", "report.docx"}
MAX_SOURCES = 60


def workspace_of(runner, job) -> Path:
    return Path(runner.store.get_project(job["project_id"])["workspace_path"])


def find_sources(root: Path, rel: str = ".") -> list[Path]:
    base = (root / rel).resolve()
    found = []
    for p in sorted(base.rglob("*")):
        if not p.is_file() or p.suffix.lower() not in DOC_TYPES:
            continue
        parts = set(p.relative_to(root).parts[:-1])
        if parts & SKIP_DIRS or p.name.lower() in OUTPUT_NAMES or p.name.lower() == "readme.md" and p.parent == root:
            continue
        found.append(p)
    return found


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:40] or "section"


READ_INSTRUCTIONS = (
    "Read {path} completely with read_document (page through it with offset until the end). For every claim relevant "
    "to the question, call add_note with the exact quote and its location (page or section). If this source "
    "disagrees with another source you've seen in the notes (search_notes), record the disagreement in the context. "
    "If the document isn't relevant to the question, add no notes and say 'not relevant' and why in your summary.")

OUTLINE_INSTRUCTIONS = (
    "Use search_notes to review all the evidence. Write outline.md for a report answering the question: a '# ' title, "
    "3 to 7 '## ' sections in a sensible order, each with bullet points that cite supporting notes like [n12]. End "
    "with a section '## Conflicts and open questions' that lists where sources disagree (cite both sides) and what the "
    "evidence doesn't answer. Don't include sources judged not relevant.")

SECTION_INSTRUCTIONS = (
    "Write the report section \"{heading}\" to {path}, following its bullets in outline.md. Start the file with "
    "'## {heading}'. Support every factual claim with note citations like [n12] (use search_notes; cite only notes "
    "that support the claim). Where sources disagree, present both and explain. Plain, precise prose; no claims "
    "without evidence.")

SUMMARY_INSTRUCTIONS = (
    "Read all section files in sections/ and write {path}: '## Summary', then a concise summary (150-300 words) "
    "answering the question directly, citing the most important notes like [n12], and noting any major disagreement "
    "between sources.")


def _initial_plan(runner, job):
    ws = workspace_of(runner, job)
    inputs = job.get("inputs") or {}
    sources = find_sources(ws, inputs.get("sources") or ".")
    if not sources:
        raise ValueError(f"No documents ({', '.join(sorted(DOC_TYPES))}) found in {ws / (inputs.get('sources') or '.')}")
    if len(sources) > MAX_SOURCES:
        raise ValueError(f"Found {len(sources)} documents; the limit is {MAX_SOURCES}. Point the job at a smaller folder.")
    tasks = [{"key": "read", "title": "Read the sources and take notes", "instructions": "-", "done_when": "-"}]
    for i, p in enumerate(sources, 1):
        rel = p.relative_to(ws).as_posix()
        tasks.append({"key": f"r{i}", "parent_key": "read", "title": f"Read {rel}",
                      "instructions": READ_INSTRUCTIONS.format(path=rel),
                      "done_when": f"Notes with verified quotes saved from {rel}, or it's reported as not relevant",
                      "checks": [{"type": "notes_for_source", "source": rel}], "params": {"source": rel}})
    tasks.append({"key": "outline", "title": "Outline the report", "instructions": OUTLINE_INSTRUCTIONS,
                  "done_when": "outline.md exists with 3-7 sections citing notes and a conflicts section",
                  "depends_on": ["read"], "review": True,
                  "checks": [{"type": "file_exists", "path": "outline.md"},
                             {"type": "file_contains", "path": "outline.md", "text": "## "},
                             {"type": "citations_valid", "path": "outline.md"}]})
    tasks.append({"key": "outline_gate", "title": "Your review of the outline", "kind": "gate", "depends_on": ["outline"],
                  "params": {"prompt": "Review outline.md in the workspace. Reply 'approve' to write the report from it, "
                                       "or describe what to change."}})
    return tasks


def _on_gate(runner, job, task, answer):
    if is_approval(answer):
        return "approve"
    outline = next(t for t in runner.jobs.list_tasks(job["id"]) if t["key"] == "outline")
    runner.jobs.add_guidance(outline["id"], f"The user reviewed your outline and asked for changes: \"{answer}\". "
                                            "Revise outline.md accordingly.")
    runner.jobs.update_task(outline["id"], status="pending", attempts=0)
    return "revise"


def _on_task_done(runner, job, task):
    if task["key"] != "outline_gate":
        return
    existing = {t["key"] for t in runner.jobs.list_tasks(job["id"])}
    if "compile" in existing:
        return
    ws = workspace_of(runner, job)
    headings = [h.strip() for h in re.findall(r"^##\s+(.+)$", (ws / "outline.md").read_text(encoding="utf-8"), re.M)]
    if not headings:
        raise ValueError("outline.md has no '## ' sections")
    new = [{"key": "write", "title": "Write the report", "instructions": "-", "done_when": "-",
            "depends_on": ["outline_gate"]}]
    section_keys = []
    for i, heading in enumerate(headings, 1):
        path = f"sections/{i:02d}-{slug(heading)}.md"
        key = f"s{i}"
        section_keys.append(key)
        new.append({"key": key, "parent_key": "write", "title": f"Write section: {heading}",
                    "instructions": SECTION_INSTRUCTIONS.format(heading=heading, path=path),
                    "done_when": f"{path} exists, starts with the heading, and cites valid notes", "review": True,
                    "checks": [{"type": "file_contains", "path": path, "text": f"## {heading}"},
                               {"type": "citations_valid", "path": path}],
                    "params": {"heading": heading, "path": path}})
    new.append({"key": "summary", "parent_key": "write", "title": "Write the summary",
                "instructions": SUMMARY_INSTRUCTIONS.format(path="sections/00-summary.md"),
                "done_when": "sections/00-summary.md exists with a cited summary", "depends_on": section_keys,
                "review": True, "checks": [{"type": "file_contains", "path": "sections/00-summary.md", "text": "## Summary"},
                                           {"type": "citations_valid", "path": "sections/00-summary.md"}]})
    new.append({"key": "compile", "title": "Compile the report", "kind": "code", "handler": "compile",
                "depends_on": ["write"], "instructions": "Assemble sections and a sources appendix.",
                "done_when": "report file exists"})
    runner.jobs.append_tasks(job["id"], new)
    runner.jobs.journal(job["id"], "plan", f"Outline approved: added {len(headings)} section tasks, a summary, and compile.")


def compile_report(runner, job, task, title_default: str = "Report", front: list[str] = (), back: list[str] = (),
                   extra_appendix: str = "") -> HandlerResult:
    ws = workspace_of(runner, job)
    inputs = job.get("inputs") or {}
    outline = ws / "outline.md"
    title_match = re.search(r"^#\s+(.+)$", outline.read_text(encoding="utf-8"), re.M) if outline.exists() else None
    title = title_match.group(1).strip() if title_match else title_default
    section_files = sorted((ws / "sections").glob("*.md")) if (ws / "sections").exists() else []
    if not section_files:
        return HandlerResult(False, "No section files found in sections/")
    ordered = [ws / f for f in front if (ws / f).exists()] + \
              [p for p in section_files if p.relative_to(ws).as_posix() not in set(front) | set(back)] + \
              [ws / f for f in back if (ws / f).exists()]
    body = "\n\n".join(p.read_text(encoding="utf-8").strip() for p in ordered)
    cited = sorted({int(m) for m in re.findall(r"\[n(\d+)\]", body)})
    notes = {n["id"]: n for n in runner.jobs.list_notes(job["id"])}
    appendix = ["## Sources and evidence"]
    for source in sorted({notes[i]["source"] for i in cited if i in notes}):
        appendix.append(f"\n### {source}")
        for i in cited:
            n = notes.get(i)
            if n and n["source"] == source:
                loc = f" ({n['location']})" if n["location"] else ""
                appendix.append(f"- **[n{i}]**{loc} {n['claim']}: “{n['quote']}”")
    report = f"# {title}\n\n{body}\n\n" + (extra_appendix.strip() + "\n\n" if extra_appendix else "") + "\n".join(appendix) + "\n"
    out = ws / (inputs.get("output") or "report.md")
    out.write_text(report, encoding="utf-8")
    written = [out.name]
    if (inputs.get("format") or "md") == "docx":
        write_docx(report, out.with_suffix(".docx"))
        written.append(out.with_suffix(".docx").name)
    return HandlerResult(True, f"Wrote {', '.join(written)}: {len(ordered)} sections, {len(cited)} cited notes "
                               f"from {len({notes[i]['source'] for i in cited if i in notes})} sources.")


def write_docx(markdown: str, path: Path) -> None:
    import docx
    d = docx.Document()
    for line in markdown.splitlines():
        stripped = line.strip()
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            d.add_heading(m.group(2), level=min(len(m.group(1)), 4))
        elif stripped.startswith(("- ", "* ")):
            d.add_paragraph(re.sub(r"\*\*(.+?)\*\*", r"\1", stripped[2:]), style="List Bullet")
        elif stripped:
            d.add_paragraph(re.sub(r"\*\*(.+?)\*\*", r"\1", stripped))
    d.save(str(path))


class ResearchReport(Template):
    def initial_plan(self, runner, job):
        return _initial_plan(runner, job)

    def on_gate(self, runner, job, task, answer):
        return _on_gate(runner, job, task, answer)

    def on_task_done(self, runner, job, task):
        _on_task_done(runner, job, task)


RESEARCH_REPORT = register(ResearchReport(
    name="research_report",
    label="Research report",
    description="Reads the documents in a folder, takes notes with verified quotes, outlines a report for your "
                "approval, writes and reviews each section, and compiles the report with a sources appendix.",
    inputs_schema={"sources": {"type": "string", "label": "Folder with the documents", "default": "."},
                   "output": {"type": "string", "label": "Report file name", "default": "report.md"},
                   "format": {"enum": ["md", "docx"], "label": "Format", "default": "md"}},
    handlers={"compile": lambda runner, job, task: compile_report(runner, job, task)},
))
