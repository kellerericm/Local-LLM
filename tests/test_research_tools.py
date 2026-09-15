import json

import docx
import pytest

from localagent.jobs.documents import closest_snippet, extract, find_quote
from test_jobs import GOOD_PLAN, call  # noqa: F401  (fixtures via conftest; helpers reused)


@pytest.fixture
def env_factory(store, settings, workspace):
    from test_jobs import Env
    return lambda responses, **kw: Env(store, settings, workspace, responses, **kw)


def make_pdf(path, pages):
    """Minimal text PDF without extra dependencies."""
    objs, kids = [], []
    font_id = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        page_id, content_id = 3 + 2 * i, 4 + 2 * i
        kids.append(f"{page_id} 0 R")
        esc = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({esc}) Tj ET".encode()
        objs.append((page_id, f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
                              f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()))
        objs.append((content_id, b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"))
    objs = [(1, b"<< /Type /Catalog /Pages 2 0 R >>"),
            (2, f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode())] + objs + \
           [(font_id, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")]
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for num, body in sorted(objs):
        offsets[num] = len(out)
        out += f"{num} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for num in range(1, len(objs) + 1):
        out += f"{offsets[num]:010d} 00000 n \n".encode()
    out += f"trailer << /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    path.write_bytes(bytes(out))


def test_extract_pdf_docx_and_txt(tmp_path):
    pdf = tmp_path / "paper.pdf"
    make_pdf(pdf, ["Sleep consolidates memories via hippocampal replay.", "Transformers lack persistent memory."])
    doc = extract(pdf, tmp_path / "cache")
    assert doc.pages == 2 and "[page 2]" in doc.text and "hippocampal replay" in doc.text
    assert extract(pdf, tmp_path / "cache").text == doc.text          # served from cache

    d = docx.Document()
    d.add_heading("Methods", level=2)
    d.add_paragraph("We trained agents with episodic memory buffers.")
    docx_path = tmp_path / "notes.docx"
    d.save(docx_path)
    assert "## Methods" in extract(docx_path).text

    txt = tmp_path / "a.md"
    txt.write_text("Working memory holds about four items.")
    assert find_quote(extract(txt).text, "working  memory holds\nabout four items")


def test_find_quote_tolerates_pdf_hyphenation():
    text = "long-term poten-\ntiation strengthens synapses"
    assert find_quote(text, "long-term potentiation strengthens synapses")
    assert not find_quote(text, "short-term depression weakens synapses")
    near = closest_snippet("filler words here " * 30 + text + " more filler text" * 30, "tiation strengthens the synapses")
    assert near is not None and "strengthens synapses" in near


def test_add_note_verifies_quotes_and_search_finds_them(env_factory, workspace):
    (workspace / "papers").mkdir()
    make_pdf(workspace / "papers" / "replay.pdf", ["Hippocampal replay during sleep supports memory consolidation."])
    (workspace / "survey.md").write_text("Retrieval-augmented models store memories outside the network weights.")
    env = env_factory([
        call("add_note", claim="Replay consolidates memory", quote="replay during sleep supports memory",
             source="papers/replay.pdf", location="p. 1"),
        call("add_note", claim="Made up", quote="neurons dream in color every night", source="papers/replay.pdf"),
        call("add_note", claim="RAG keeps memory external", quote="store memories outside the network weights",
             source="survey.md"),
        call("search_notes", query="memory consolidation sleep"),
        call("fail_task", reason="test done"),
    ])
    job = env.job()
    env.plan(job["id"], plan=GOOD_PLAN[:1])
    env.runner._tick()
    notes = env.jobs.list_notes(job["id"])
    assert [n["source"] for n in notes] == ["papers/replay.pdf", "survey.md"]
    run = env.jobs.list_runs(job["id"])[-1]
    results = [m["content"] for m in env.jobs.list_run_messages(run["id"]) if m["role"] == "tool"]
    assert results[0].startswith("Saved note n") and "doesn't appear word for word" in results[1]
    assert "[n" in results[3] and "replay" in results[3].lower()
    assert env.jobs.search_notes(job["id"], "weights")[0]["source"] == "survey.md"
    env.jobs.delete_job(job["id"])
    assert env.jobs.search_notes(job["id"], "weights") == []            # FTS cleaned up via cascade + trigger
